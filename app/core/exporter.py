import os
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from typing import Dict, Any

class ExportManager:
    """Exports approved lesson plans, worksheets, and parent reports to PDF, DOCX, and WhatsApp."""

    @staticmethod
    def export_to_docx(title: str, content: Dict[str, Any], filepath: str) -> str:
        os.makedirs(os.path.dirname(filepath), exist_ok=True)
        doc = Document()
        doc.add_heading(title, 0)
        
        for key, val in content.items():
            doc.add_heading(key.replace("_", " ").title(), level=1)
            if isinstance(val, list):
                for item in val:
                    if isinstance(item, dict):
                        doc.add_paragraph(" • " + ", ".join(f"{k}: {v}" for k, v in item.items()))
                    else:
                        doc.add_paragraph(f" • {item}")
            elif isinstance(val, dict):
                for sub_k, sub_v in val.items():
                    doc.add_paragraph(f"{sub_k}: {sub_v}")
            else:
                doc.add_paragraph(str(val))
                
        doc.save(filepath)
        return filepath

    @staticmethod
    def export_to_pdf(title: str, content: Dict[str, Any], filepath: str) -> str:
        os.makedirs(os.path.dirname(filepath), exist_ok=True)
        c = canvas.Canvas(filepath, pagesize=letter)
        y = 750
        c.setFont("Helvetica-Bold", 16)
        c.drawString(50, y, title[:60])
        y -= 30

        c.setFont("Helvetica", 10)
        for key, val in content.items():
            if y < 80:
                c.showPage()
                y = 750
                c.setFont("Helvetica", 10)

            c.setFont("Helvetica-Bold", 11)
            c.drawString(50, y, f"{key.replace('_', ' ').title()}:")
            y -= 15
            c.setFont("Helvetica", 9)
            
            val_str = str(val)
            lines = [val_str[i:i+80] for i in range(0, len(val_str), 80)]
            for line in lines[:4]:
                c.drawString(60, y, line)
                y -= 12
            y -= 8

        c.save()
        return filepath

    @staticmethod
    def format_whatsapp_message(parent_name: str, student_name: str, report: Dict[str, Any]) -> str:
        strengths = report.get("strengths", "Strong participation and curiosity.")
        gaps = report.get("learning_gaps", "Needs minor review on core logic.")
        home_action = report.get("actionable_home_steps", "A quick 5-min recap.")
        
        return (
            f"🌟 *Classroom Progress Update for {student_name}*\n\n"
            f"Dear {parent_name},\n"
            f"Here is a summary of today's learning:\n\n"
            f"✨ *Key Strengths*: {strengths}\n"
            f"🎯 *Focus Area*: {gaps}\n"
            f"💡 *Fun 5-Min Home Activity*: {home_action}\n\n"
            f"Thank you for your continuous support! 📚"
        )

export_manager = ExportManager()
